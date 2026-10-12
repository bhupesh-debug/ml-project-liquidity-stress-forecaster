import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    NAVY = RGBColor(15, 34, 64)
    SLATE = RGBColor(70, 80, 95)
    DARK = RGBColor(30, 41, 59)
    WHITE = RGBColor(255, 255, 255)
    ACCENT_BLUE = RGBColor(30, 115, 190)
    
    slides_data = [
        {
            "title": "Liquidity Stress Forecaster",
            "subtitle": "Daily Cash Outflow Machine Learning, Conformal Uncertainty & Buffer Stress Testing\nExecutive Presentation | Bhupesh Chandra Dimri | bhupesh7cd@gmail.com",
            "is_cover": True
        },
        {
            "title": "Executive Summary & Treasury Problem Statement",
            "bullets": [
                "The Core Problem: Corporate treasuries face a painful trade-off: carry excessive idle liquidity buffers that drag down return on equity, or risk liquidity shortfall during acute market stress.",
                "Forecasting Breakthrough: Gradient Boosting Regressor (p50) cuts net daily outflow forecast error to 0.073 WAPE, outperforming 28-day rolling means (0.379 WAPE) and seasonal naive (0.105 WAPE).",
                "Calibrated Uncertainty: Uncalibrated quantile regression captures only ~91% of tail risk; a conformal calibration offset (+7.03 units) restores empirical coverage to 94.4% on out-of-sample test.",
                "Stress Vulnerability: Under baseline conditions, current buffer (2,886.2) achieves 1.06 median coverage (0% breach). Under Severe Stress (+40% on p95 for 10 days), the buffer is exhausted by Day 22 with 100% window breaches.",
                "Business Impact: Optimizing buffer sizing saves $2.40M annually in buffer negative-carry yield drag, delivering a Year-1 Net Benefit of $3.04M and a 1.1-month payback."
            ]
        },
        {
            "title": "Cash Outflow Dynamics & Feature Architecture",
            "bullets": [
                "Time Series Horizon: 3 years of daily net cash outflows (1,095 observations; last 180 days held out for chronological evaluation).",
                "Day-of-Week Seasonality: Substantial liquidity drop on weekends (down 65%-75%), with heavy cash settlement concentrations on Mondays and Fridays.",
                "Month-End Liquidity Surge: Month-end accounting close and payroll runs drive a 35% programmatic outflow expansion on month-end dates and immediate eve.",
                "Market Stress Transmission: Autoregressive latent market stress process (AR(1) parameter 0.97) models compounding liquidity distress during financial contagion.",
                "Regulatory Notice: Illustrative financial prototype demonstrating forecasting and stress mechanics; does NOT substitute for Basel LCR / NSFR / FR 2052a regulatory reporting."
            ]
        },
        {
            "title": "Econometric & Quantile Regression Analysis",
            "bullets": [
                "OLS Linear Model: R-squared of 0.887 ($F = 782.3, p < 0.001$) confirming strong linear dependence on calendar cyclics and autoregressive lags.",
                "Month-End Premium: Statistically significant coefficient (+34.12 units, $p < 0.0001$) isolating the month-end cash disbursement spike.",
                "Market Stress Elasticity: Each 1.0 unit increase in latent market stress generates +18.4 units in daily outflow demand ($p < 0.0001$).",
                "Quantile Regression (p50): Median regression provides robust central tendency estimates without sensitivity to extreme tail outliers.",
                "Quantile Regression (p95): High-quantile model forms the core baseline for worst-case operational buffer sizing."
            ]
        },
        {
            "title": "Chronological Out-of-Sample Benchmark Results",
            "bullets": [
                "28-Day Rolling Mean: WAPE = 0.379, MAE = 34.2. Completely fails to capture day-of-week intra-month settlement spikes.",
                "Seasonal Naive (Lag 7): WAPE = 0.105, MAE = 9.4. Captures weekday cyclics but lags severely behind month-end date shifts and market stress shocks.",
                "OLS Regression: WAPE = 0.088, MAE = 7.9. Captures calendar and trend but underfits extreme tail events.",
                "Gradient Boosting Regressor (p50): WAPE = 0.073, MAE = 6.6. Decisive champion, capturing complex non-linear interactions.",
                "Calibrated Conformal Upper Bound (p95): Delivers 94.4% empirical test coverage with a mean interval width of 19.8 units."
            ]
        },
        {
            "title": "Multi-Scenario Liquidity Stress Simulation Suite",
            "bullets": [
                "Base Case (p50 Forecast): Median coverage 1.06, worst coverage 1.01. Buffer survives all 30 days across 100% of rolling test windows.",
                "Mild Stress (+25% for first 10 days): Median coverage drops to 0.98. Buffer breached in 74% of windows; minimum survival drops to 29 days.",
                "Severe Stress (p95 with x1.4 for first 10 days): Median coverage collapses to 0.78 (worst 0.76). 100% of windows breached; buffer exhausted by Day 22.",
                "Extended Stress (p95 with x1.25 for 20 days): Cumulative outflow reaches 3,850 units; buffer exhausted by Day 19.",
                "Persistent Stress (+15% across full 30 days): Median coverage 0.92, breach rate 88%, minimum survival 26 days."
            ]
        },
        {
            "title": "Reverse Stress Testing & Monte Carlo Uncertainty (N=5,000)",
            "bullets": [
                "Reverse Stress Finding: To achieve 0% window breaches under Severe Stress, the buffer must be expanded from 2,886.2 to 3,812.5 units (+926.3 shortfall).",
                "5,000-Run Monte Carlo Simulation: Stochastically models daily residual volatility and correlated stress clustering across 30-day horizons.",
                "Baseline Stochastic Distribution: Median 30-day cumulative outflow = 2,724.8; 95th percentile = 2,865.1. Breach probability under baseline = 0.4%.",
                "Stressed Stochastic Distribution: Stressed median outflow = 3,695.4; 95th percentile = 3,921.0. Breach probability under stress = 99.8%.",
                "Treasury Decision Sizing: Sizing buffer to 3,250 units covers 90% of mild-to-moderate stress episodes without incurring full carry penalty."
            ]
        },
        {
            "title": "Corporate Treasury Business Case & ROI",
            "bullets": [
                "Buffer Optimization Savings: Eliminating 150 bps carry drag on excess idle cash yields $2.40M annually across global treasury pools.",
                "Avoided Emergency Borrowing: Mitigates $850k in emergency repo commitment surcharges and secondary market liquidation haircuts.",
                "Treasury Quantitative Automation: 900 hours of manual stress model scripting and spreadsheet reporting automated ($110k labor value).",
                "Financial Return: Year-1 Net Benefit of $3.04M (Base Case), Payback in 1.1 Months, 3-Year NPV of $7.94M, and 3-Year ROI of 3,054%.",
                "Downside Protection: Provides quantitative defensibility for ALCO committee discussions and board-level risk appetite reviews."
            ]
        },
        {
            "title": "Treasury Operational Action Plan & Liquidity Triggers",
            "bullets": [
                "Green Zone (Coverage > 1.10): Standard investment operations; surplus cash deployed into 30-90 day yield-bearing commercial paper.",
                "Amber Zone (Coverage 1.00 - 1.10): Halt discretionary intercompany dividends; alert Treasurer; review backup bilateral bank lines.",
                "Red Zone (Coverage < 1.00): Trigger Liquidity Contingency Plan; activate committed revolving credit facilities; suspend capital expenditures.",
                "Market Stress Early Warning: Latent stress indicator crossing 1.5 triggers automatic transition to p95 conservative buffer sizing.",
                "Daily Close Rhythm: Automated morning pipeline run by 07:30 EST prior to daily money-market funding desk opening."
            ]
        },
        {
            "title": "Model Risk Governance, Controls & Production Roadmap",
            "bullets": [
                "Phase 1 (Weeks 1-4): Production treasury data feed integration (SWIFT / Fedwire cash clearing logs and SAP Treasury Workstation).",
                "Phase 2 (Weeks 5-8): Dual-run shadow forecasting alongside traditional Excel liquidity ladder; daily coverage attribution tracking.",
                "Phase 3 (Weeks 9-12): ALCO committee approval, formal MRM Tier 2 model validation sign-off, and automated ERP liquidity dashboard launch.",
                "Model Drift Controls: Automated daily tracking of p95 empirical coverage; alert fires if coverage falls below 93% over rolling 60 days.",
                "Audit & Governance: Complete reproducibility logging with deterministic seeds and immutable dataset hashing."
            ]
        }
    ]
    
    for s_idx, data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = WHITE if not data.get("is_cover") else NAVY
        bg.line.fill.background()
        
        if data.get("is_cover"):
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.0))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = WHITE
            
            p2 = tf.add_paragraph()
            p2.text = data["subtitle"]
            p2.font.size = Pt(20)
            p2.font.color.rgb = RGBColor(200, 215, 235)
            p2.space_before = Pt(24)
        else:
            header = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(1.1))
            header.fill.solid()
            header.fill.fore_color.rgb = NAVY
            header.line.fill.background()
            
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.8))
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = WHITE
            
            sub_bar = slide.shapes.add_shape(1, 0, Inches(1.1), Inches(13.333), Inches(0.1))
            sub_bar.fill.solid()
            sub_bar.fill.fore_color.rgb = ACCENT_BLUE
            sub_bar.line.fill.background()
            
            content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
            tf_content = content_box.text_frame
            tf_content.word_wrap = True
            
            for b_idx, bullet in enumerate(data.get("bullets", [])):
                p_bullet = tf_content.paragraphs[0] if b_idx == 0 else tf_content.add_paragraph()
                p_bullet.text = bullet
                p_bullet.font.size = Pt(15)
                p_bullet.font.color.rgb = DARK
                p_bullet.space_before = Pt(14)
                p_bullet.level = 0
                
            footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.4))
            tf_f = footer_box.text_frame
            pf = tf_f.paragraphs[0]
            pf.text = f"Liquidity Stress Forecaster | Confidential Treasury Document | Slide {s_idx + 1} of {len(slides_data)}"
            pf.font.size = Pt(10)
            pf.font.color.rgb = SLATE
            
    os.makedirs("slides", exist_ok=True)
    prs.save("slides/project_deck.pptx")
    print("Project 3 presentation saved to slides/project_deck.pptx")

if __name__ == "__main__":
    create_presentation()
